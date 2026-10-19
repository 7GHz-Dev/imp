# -*- coding: utf-8 -*-
"""
อ่านไฟล์ invoice (อ่านอย่างเดียว ไม่แก้ไฟล์)

อ่าน "ค่าที่ Excel คำนวณไว้" (data_only) ของตารางสรุป และหาช่องเลขเครื่อง
ตาม template  ไฟล์ต้องเคยเปิดและ Save ด้วย Excel มาแล้ว ค่าสูตรถึงจะมี
"""
import re
from dataclasses import dataclass, field
from decimal import Decimal

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string

from . import layouts as L
from .values import clean, to_decimal, fmt_date

# เลขเครื่องที่สุ่มแล้วมีรูปแบบ  <หัวเครื่อง>-<เลข 7 หลัก>  เช่น HR15-1234567
FULL_ENGINE = re.compile(r'^(.*\S)-(\d{7})$')


@dataclass
class Car:
    row: int
    car_no: int
    brand: str
    model: str
    chassis: str
    weight: Decimal
    year: str
    cc: int
    fuel: str
    price: Decimal
    tariff_cell: str
    desc1: str
    desc2: str
    use_chassis: str
    use_weight: Decimal
    use_price: Decimal


@dataclass
class EngineCell:
    """ช่องเลขเครื่องในส่วน invoice ของผู้ขาย (ที่จะเขียนเลขสุ่มลงไป)"""
    row: int
    col: int
    chassis: str
    raw: str        # ข้อความเดิมในช่อง เช่น "HR15" หรือ "LDA-MF6 (GASOLINE/1330)"
    head: str       # ส่วนเลขเครื่อง เช่น "HR15" / "LDA-MF6"

    @property
    def existing_number(self):
        """ถ้าช่องนี้มีเลข 7 หลักต่อท้ายแล้ว คืน (หัวเครื่อง, เลข)  ไม่มีคืน None"""
        m = FULL_ENGINE.match(self.head)
        return (m.group(1), m.group(2)) if m else None

    def value_with(self, full_engine):
        """ข้อความที่จะเขียนกลับลงช่อง เมื่อได้เลขเครื่องเต็มแล้ว"""
        return self.raw.replace(self.head, full_engine, 1)


@dataclass
class Invoice:
    path: str
    file_type: str = ''
    invoice_no: str = ''
    invoice_date: str = ''
    bl_no: str = ''
    cars: list = field(default_factory=list)
    engines: list = field(default_factory=list)
    tariff_sheet: object = None
    errors: list = field(default_factory=list)
    warnings: list = field(default_factory=list)


def _label(v):
    return re.sub(r'[\s.:]', '', clean(v).upper())


def read_invoice(path):
    inv = Invoice(path=path)
    try:
        wb = load_workbook(path, data_only=True)
    except Exception as e:
        inv.errors.append('เปิดไฟล์ Excel ไม่ได้: %s (ต้องเป็น .xlsx)' % e)
        return inv
    ws = wb.worksheets[0]

    inv.file_type = L.normalize_type(ws[L.TYPE_CELL].value)
    inv.invoice_date = fmt_date(ws[L.DATE_CELL].value)
    inv.invoice_no = clean(ws[L.INVOICE_CELL].value)
    inv.bl_no = clean(ws[L.BL_CELL].value)
    if 'datatariff' in wb.sheetnames:
        inv.tariff_sheet = wb['datatariff']

    if inv.file_type not in L.LAYOUTS:
        inv.errors.append('ไม่รู้จักชนิดไฟล์ %r ในช่อง %s (รองรับ: %s)'
                          % (inv.file_type, L.TYPE_CELL, ', '.join(L.LAYOUTS)))
        return inv
    for label, value in (('เลข Invoice', inv.invoice_no), ('เลข BL', inv.bl_no),
                         ('วันที่ Invoice', inv.invoice_date)):
        if not value:
            inv.errors.append('ไม่มี%s ในตารางสรุป (ช่อง P4-P14)' % label)

    inv.cars = _read_cars(ws, inv.errors)
    if not inv.cars and not inv.errors:
        inv.errors.append('ไม่พบรถในตารางสรุป (แถว 20 ลงไป คอลัมน์ W:AF) — '
                          'ถ้าเพิ่งกรอกข้อมูล ให้เปิดไฟล์ใน Excel กด Save แล้วอัปโหลดใหม่')

    layout = L.LAYOUTS[inv.file_type]
    if isinstance(layout, L.ColumnLayout):
        inv.engines = _engines_by_column(ws, layout, inv.errors)
    else:
        inv.engines = _engines_by_block(ws, layout, inv.errors)
    return inv


def _read_cars(ws, errors):
    col = {k: column_index_from_string(v) for k, v in L.SUMMARY_COLS.items()}
    cars = []
    for r in range(L.FIRST_CAR_ROW, ws.max_row + 1):
        get = lambda k: ws.cell(r, col[k]).value
        chassis = clean(get('chassis'))
        if chassis in ('', '0'):      # แถวว่างที่สูตร VLOOKUP คืน 0
            continue
        nums = {k: to_decimal(get(k))
                for k in ('car_no', 'weight', 'year', 'cc', 'price', 'use_weight', 'use_price')}
        missing = [k for k, v in nums.items() if v is None]
        missing += [k for k in ('brand', 'fuel', 'desc1', 'desc2', 'use_chassis') if not clean(get(k))]
        if missing:
            errors.append('แถว %d (%s): ข้อมูลไม่ครบหรือไม่ใช่ตัวเลข: %s'
                          % (r, chassis, ', '.join(missing)))
            continue
        cars.append(Car(
            row=r,
            car_no=int(nums['car_no']),
            brand=clean(get('brand')),
            model=clean(get('model')),
            chassis=chassis,
            weight=nums['weight'],
            year=str(int(nums['year'])),
            cc=int(nums['cc']),
            fuel=clean(get('fuel')),
            price=nums['price'],
            tariff_cell=clean(get('tariff')),
            desc1=clean(get('desc1')),
            desc2=clean(get('desc2')),
            use_chassis=clean(get('use_chassis')),
            use_weight=nums['use_weight'],
            use_price=nums['use_price'],
        ))
    return cars


def _engines_by_column(ws, lay, errors):
    c_col = column_index_from_string(lay.chassis_col)
    e_col = column_index_from_string(lay.engine_col)
    h_chassis = _label(ws.cell(lay.header_row, c_col).value)
    h_engine = _label(ws.cell(lay.header_row, e_col).value)
    if 'CHASSIS' not in h_chassis or 'ENGINE' not in h_engine:
        errors.append('หัวตารางแถว %d ไม่ตรง template (คาดว่า %s=CHASSIS, %s=ENGINE แต่พบ %r, %r)'
                      % (lay.header_row, lay.chassis_col, lay.engine_col, h_chassis, h_engine))
        return []
    out = []
    for r in range(lay.header_row + 1, ws.max_row + 1):
        chassis = clean(ws.cell(r, c_col).value)
        if not chassis:
            continue
        raw = clean(ws.cell(r, e_col).value)
        if not raw:
            errors.append('แถว %d (%s): ไม่มีเลขเครื่องในคอลัมน์ %s' % (r, chassis, lay.engine_col))
            continue
        out.append(EngineCell(row=r, col=e_col, chassis=chassis, raw=raw, head=raw))
    return out


def _engines_by_block(ws, lay, errors):
    l_col = column_index_from_string(lay.label_col)
    v_col = column_index_from_string(lay.value_col)
    out = []
    chassis = None
    for r in range(lay.first_row, ws.max_row + 1):
        label = _label(ws.cell(r, l_col).value)
        value = ws.cell(r, v_col).value
        if label == lay.chassis_label:
            chassis = clean(value) or None
        elif label == lay.engine_label:
            raw = ' '.join(str(value or '').split())
            head = raw.split('(', 1)[0].strip()
            if chassis is None:
                errors.append('แถว %d: พบเลขเครื่องแต่ไม่มี CHASSIS NO. ก่อนหน้า' % r)
            elif not head:
                errors.append('แถว %d (%s): ไม่มีเลขเครื่องในคอลัมน์ %s' % (r, chassis, lay.value_col))
            else:
                out.append(EngineCell(row=r, col=v_col, chassis=chassis, raw=raw, head=head))
            chassis = None
    return out
