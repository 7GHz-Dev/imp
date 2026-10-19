# -*- coding: utf-8 -*-
"""
Mapping ข้อมูลรถ -> ไฟล์ CSV สำหรับนำเข้า NETBAY

รูปแบบไฟล์เหมือนโปรแกรมเดิม (adjust_invoice + pandas.to_csv) ทุกอย่าง:
  บรรทัด 1  หัวคอลัมน์จาก data/csv_template.xlsx (shippingMark ตัวที่ 2 เป็น shippingMark.1)
  บรรทัด 2  "//" ตามด้วย N/A
  บรรทัด 3+ รถ 1 คันต่อบรรทัด  ค่าคงที่ (610, MAESOT FREE ZONE, ...) มาจากแถว 3 ของ template
  ช่องว่าง = N/A   encoding TIS-620
แก้ค่าคงที่ (เช่นเลขใบอนุญาต FX_permitNo1) ได้ที่แถว 3 ของ csv_template.xlsx
"""
import csv
import io

from openpyxl import load_workbook

from .values import clean, fmt_number

THAI_CAR = 'รถยนต์เก่าใช้แล้วนำเข้ามาเพื่อปรับสภาพและส่งออก'
THAI_TRUCK = 'รถบรรทุกเก่าใช้แล้วนำเข้ามาเพื่อปรับสภาพและส่งออก'
NA = 'N/A'


class CsvTemplate:
    def __init__(self, path):
        ws = load_workbook(path).worksheets[0]
        header = [clean(c.value) for c in ws[1]]
        while header and not header[-1]:
            header.pop()
        self.columns = _dedupe(header)
        self.defaults = [clean(ws.cell(3, i + 1).value) or NA for i in range(len(header))]


def _dedupe(names):
    """ชื่อซ้ำ -> ชื่อ.1 แบบเดียวกับที่ pandas ทำในโปรแกรมเดิม"""
    seen, out = {}, []
    for n in names:
        if n in seen:
            seen[n] += 1
            out.append('%s.%d' % (n, seen[n]))
        else:
            seen[n] = 0
            out.append(n)
    return out


def load_brand_country(path):
    """ORIGINAL NETBAY.xlsx: แถว 2 = รหัสประเทศ, ใต้ลงไป = ยี่ห้อรถของประเทศนั้น  (ไม่อยู่ในรายการ = JP)"""
    ws = load_workbook(path, data_only=True, read_only=True).worksheets[0]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    out = {}
    for j, country in enumerate(rows[0] if rows else []):
        for r in rows[1:]:
            if j < len(r) and clean(r[j]):
                out[clean(r[j]).upper()] = clean(country)
    return out


def split_engine(desc2):
    """
    desc2 ที่มีเลขสุ่มแล้ว เช่น "HR15-1234567(GASOLINE/1490),2011"
    คืน (endesc2, productAttribute2) = ("HR15(GASOLINE/1490),2011", "HR15-1234567")
    ตรรกะเดียวกับ adjust_invoice เดิม: ตัดตั้งแต่ขีดตัวสุดท้ายก่อนวงเล็บ
    """
    p = desc2.rfind('(')
    dash = desc2[:p].rfind('-')
    return desc2[:dash] + desc2[p:], desc2.split('(')[0]


def add_number(desc2, number):
    """ต่อเลขสุ่มท้ายเลขเครื่องใน desc2  "HR15(GASOLINE/1490),2011" -> "HR15-1234567(GASOLINE/1490),2011"
    (ได้ค่าเดียวกับที่สูตรใน Excel คำนวณหลังเติมเลขลงช่องเลขเครื่อง)"""
    p = desc2.find('(')
    return '%s-%s%s' % (desc2[:p], number, desc2[p:])


def map_car(car, invoice, tariff_code, desc2_full, brand_country):
    """รถ 1 คัน -> {ชื่อคอลัมน์ CSV: ค่า}  เฉพาะช่องที่ไม่ใช่ค่าคงที่"""
    endesc2, engine = split_engine(desc2_full)
    v = {
        'invoiceNo': invoice.invoice_no,
        'invoiceDate': invoice.invoice_date,
        'invoiceItem': str(car.car_no),
        'productCode': '%dCC' % car.cc,
        'ins_tariffCode': tariff_code,
        'englishDescriptionOfGoods': car.desc1.replace(' ,', ','),
        'endesc2': endesc2,
        'brandName': car.brand,
        'customsProductCode': car.fuel,
        'productYear': car.year,
        'productAttribute1': car.use_chassis,
        'productAttribute2': engine,
        'netWeight': fmt_number(car.use_weight),
        'itemGrossWeight': fmt_number(car.use_weight),
        'invoiceAmountForeign': fmt_number(car.use_price),
    }
    heading = tariff_code[4:8]
    if heading == '8703':
        v['thaiDescriptionOfGoods'] = THAI_CAR
    elif heading == '8704':
        v['thaiDescriptionOfGoods'] = THAI_TRUCK
        if car.use_weight > 5000:
            v['ins_privilegeCode'] = '999'
            v['ins_tariffSequence'] = '65001'
    country = brand_country.get(car.brand.upper())
    if country:
        v['originCountry'] = country
    return v


def build_rows(template, mapped):
    """mapped = list ของ dict จาก map_car -> list ของแถว (เรียงตามคอลัมน์ template)"""
    index = {name: i for i, name in enumerate(template.columns)}
    rows = []
    for m in mapped:
        row = list(template.defaults)
        for name, value in m.items():
            row[index[name]] = value if value != '' else NA
        rows.append(row)
    return rows


def render(template, rows, encoding='tis-620'):
    """แถวทั้งหมด -> bytes ของไฟล์ CSV  (ตรวจว่าทุกตัวอักษรเข้ารหัสได้)"""
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator='\r\n')
    w.writerow(template.columns)
    w.writerow(['//'] + [NA] * (len(template.columns) - 1))
    w.writerows(rows)
    text = buf.getvalue()
    try:
        return text.encode(encoding)
    except UnicodeEncodeError as e:
        bad = text[e.start:e.end]
        line = text[:e.start].count('\n') + 1
        raise ValueError('ตัวอักษร %r บรรทัด %d เข้ารหัส %s ไม่ได้' % (bad, line, encoding)) from None


def totals(template, rows):
    """ยอดรวมสำหรับกรอกหัวใบขน NETBAY (คำนวณจากแถว CSV แบบโปรแกรมเดิม)"""
    from decimal import Decimal
    col = {name: i for i, name in enumerate(template.columns)}
    s = lambda name: sum(Decimal(r[col[name]]) for r in rows)
    return {
        'package_amount': fmt_number(s('packageAmount')),
        'package_unit': rows[0][col['packageUnit']],
        'gross_weight': fmt_number(s('itemGrossWeight')),
        'price': fmt_number(s('invoiceAmountForeign')),
        'currency': rows[0][col['currency']],
    }
