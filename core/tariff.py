# -*- coding: utf-8 -*-
"""
หาพิกัดศุลกากรจากตาราง datatariff (ยกตรรกะมาจาก excel_to_csv ของโปรแกรมเดิม)

จับคู่ ยี่ห้อ+รุ่น+เชื้อเพลิง (ตัดช่องว่างและขีดออก) และ CC/น้ำหนัก อยู่ในช่วง
  เจอ 1 พิกัด  -> '0000' + พิกัด
  เจอหลายพิกัด -> '0000XXXX' (ต้องตรวจเอง)  ถ้าทุกแถวที่เจอเป็นพิกัดเดียวกันจะใช้พิกัดนั้น
  ไม่เจอ       -> ใช้ค่าในคอลัมน์ tariff ของ invoice ตามเดิม
"""
from dataclasses import dataclass

from openpyxl import load_workbook

from .values import clean, to_decimal

AMBIGUOUS = '0000XXXX'


@dataclass(frozen=True)
class TariffRow:
    num: str
    code: str
    key: str
    cc_start: int
    cc_end: int
    w_start: int
    w_end: int


def _key(*parts):
    return ''.join(clean(p) for p in parts).replace(' ', '').replace('-', '').upper()


def load_rows(ws):
    """อ่านตาราง: หัวตารางแถว 2  คอลัมน์ A:I = trf_num, trf_code, brand, model, fuel, cc_start, cc_end, w_start, w_end"""
    rows = []
    for v in ws.iter_rows(min_row=3, max_col=9, values_only=True):
        num, code, brand, model, fuel = v[:5]
        rng = [to_decimal(x) for x in v[5:9]]
        if not clean(code) or not clean(brand) or None in rng:
            continue
        rows.append(TariffRow(clean(num), clean(code), _key(brand, model, fuel), *[int(x) for x in rng]))
    return rows


def load_file(path):
    return load_rows(load_workbook(path, data_only=True, read_only=True).worksheets[0])


def lookup(rows, car):
    """คืน (พิกัด, หมายเหตุ)  หมายเหตุว่าง = เจอพิกัดเดียว"""
    key = _key(car.brand, car.model, car.fuel)
    hits = [t for t in rows
            if t.key == key and t.cc_start <= car.cc <= t.cc_end and t.w_start <= car.weight <= t.w_end]
    codes = sorted({t.code for t in hits})
    if len(codes) == 1:
        return '0000' + codes[0], ''
    if len(codes) > 1:
        return AMBIGUOUS, 'เจอหลายพิกัด (%s) ลำดับ %s' % (', '.join(codes), ', '.join(t.num for t in hits))
    return car.tariff_cell or 'XXXX', 'ไม่พบพิกัดสำหรับ %s %s %s %dCC %sKG' % (
        car.brand, car.model, car.fuel, car.cc, car.weight)
