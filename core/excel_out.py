# -*- coding: utf-8 -*-
"""
เขียนเลขเครื่องกลับลงสำเนาไฟล์ invoice

Excel (Windows ที่มี Microsoft Excel) : เปิดด้วย Excel จริง — รูปแบบ/รูปภาพครบ สูตรคำนวณใหม่ทันที
openpyxl (สำรอง)                      : สูตรยังอยู่ Excel จะคำนวณใหม่ตอนเปิดไฟล์
                                       แต่รูปภาพ/รูปแบบพิเศษบางอย่างอาจหาย
"""
import os
import shutil

from openpyxl import load_workbook


def excel_available():
    try:
        import win32com.client  # noqa: F401
        return True
    except ImportError:
        return False


def write_engines(src, dst, updates, check_rows=(), method='auto'):
    """
    updates    = [(row, col, value), ...]  แบบ 1-based บนชีตแรก
    check_rows = แถวตารางสรุปที่อยากอ่านคอลัมน์ Q กลับมาหลัง Excel คำนวณ (ไว้ตรวจ)
    คืน (วิธีที่ใช้, {row: ค่า Q หลังคำนวณ})  — openpyxl คืน {} เพราะคำนวณสูตรไม่ได้
    """
    if method in ('auto', 'excel') and excel_available():
        try:
            return 'excel', _write_with_excel(src, dst, updates, check_rows)
        except Exception:
            if method == 'excel':
                raise
    _write_with_openpyxl(src, dst, updates)
    return 'openpyxl', {}


def _write_with_excel(src, dst, updates, check_rows):
    import pythoncom
    import win32com.client

    shutil.copy2(src, dst)
    pythoncom.CoInitialize()
    excel = wb = ws = None
    try:
        excel = win32com.client.DispatchEx('Excel.Application')   # instance ใหม่ ไม่ยุ่งกับ Excel ที่ผู้ใช้เปิดอยู่
        excel.Visible = False
        excel.DisplayAlerts = False
        wb = excel.Workbooks.Open(os.path.abspath(dst), UpdateLinks=0)
        try:
            ws = wb.Worksheets(1)
            for row, col, value in updates:
                ws.Cells(row, col).Value = value
            excel.Calculate()
            after = {r: ws.Range('Q%d' % r).Value for r in check_rows}
            wb.Save()
        finally:
            wb.Close(SaveChanges=False)
        return after
    finally:
        ws = wb = None              # ปล่อย object ของ workbook ก่อนปิด Excel
        if excel is not None:
            excel.Quit()
            excel = None
        pythoncom.CoUninitialize()


def _write_with_openpyxl(src, dst, updates):
    wb = load_workbook(src)
    ws = wb.worksheets[0]
    for row, col, value in updates:
        ws.cell(row=row, column=col).value = value
    wb.save(dst)
