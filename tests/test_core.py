# -*- coding: utf-8 -*-
import csv
import io
import os

import pytest
from openpyxl import load_workbook

from conftest import CARS, fixture, make_jantrading, make_msfz
from core import netbay_csv as NC
from core import pipeline
from core.excel_out import excel_available
from core.values import fmt_date, fmt_number, to_decimal


def read_csv(run, settings):
    path = os.path.join(settings.output_dir, run['run_id'], run['csv'])
    with open(path, 'rb') as f:
        raw = f.read()
    rows = list(csv.reader(io.StringIO(raw.decode('tis-620'), newline='')))
    return raw, rows, [dict(zip(rows[0], r)) for r in rows[2:]]


# ---------- ค่าพื้นฐาน ----------

def test_values():
    assert fmt_number(to_decimal(1370.0)) == '1370'
    assert fmt_number(to_decimal('172700.10')) == '172700.1'
    assert fmt_date('7/8/2026') == '07/08/2026'
    assert to_decimal('#N/A') is None


def test_engine_text_helpers():
    full = NC.add_number('HR15(GASOLINE/1490),2011', '1234567')
    assert full == 'HR15-1234567(GASOLINE/1490),2011'
    assert NC.split_engine(full) == ('HR15(GASOLINE/1490),2011', 'HR15-1234567')
    assert NC.split_engine('DAA-GP2-7654321(GASOLINE/1330),2011') == ('DAA-GP2(GASOLINE/1330),2011', 'DAA-GP2-7654321')


# ---------- ฐานข้อมูลเลขเครื่อง ----------

def test_issue_never_repeats_existing_numbers(db, monkeypatch):
    with db.transaction() as tx:
        tx.add('1111111', 'X-1111111', 'OLD-1')
    seq = iter([1111111, 1111111, 2222222])
    monkeypatch.setattr('core.engine_db._rng.randint', lambda a, b: next(seq))
    with db.transaction() as tx:
        number, full = tx.issue('HR15', 'NEW-1', 'BL1', 'INV1', 'f.xlsx')
    assert (number, full) == ('2222222', 'HR15-2222222')
    assert db.find('2222222')[0]['BL_NO'] == 'BL1'


def test_transaction_rolls_back_on_error(db):
    with pytest.raises(RuntimeError):
        with db.transaction() as tx:
            tx.issue('HR15', 'C-1')
            raise RuntimeError('boom')
    assert db.count() == 0


def test_import_csv(tmp_path, db):
    p = tmp_path / 'old.csv'
    p.write_text('﻿RANDOM_ENG,FULL_ENG,CHASSIS,DATE_RAN\n1000021,1KD-1000021,KDH206-8015656,2026-05-20 17:01:07.000\n',
                 encoding='utf-8')
    assert db.import_csv(str(p)) == (1, 1)
    assert db.import_csv(str(p)) == (0, 1)   # รันซ้ำไม่เพิ่ม


# ---------- invoice -> CSV ----------

def test_msfz_end_to_end(tmp_path, settings, db):
    src = make_msfz(str(tmp_path / 'inv.xlsx'))
    a = pipeline.analyze(src, settings, db)
    assert a.ok, a.errors
    assert [l.status for l in a.lines] == ['new'] * 3
    assert db.count() == 0                       # analyze ไม่เขียนอะไร

    run = pipeline.commit(src, settings, db, 'inv.xlsx')
    assert db.count() == 3
    engines = [c['engine'] for c in run['cars']]
    assert [e.split('-')[0] for e in engines] == ['HR15', '2GR', 'M271']
    assert all(len(e.split('-')[1]) == 7 for e in engines)

    ws = load_workbook(os.path.join(settings.output_dir, run['run_id'], run['excel'])).worksheets[0]
    assert [ws['J%d' % r].value for r in (20, 21, 22)] == engines

    raw, rows, data = read_csv(run, settings)
    assert raw.count(b'\r\n') == len(rows)
    assert rows[0][31] == 'shippingMark.1' and rows[1][0] == '//'
    first = data[0]
    assert first['//Document Type'] == '610'
    assert first['companyName'] == 'บริษัท เเม่สอด ฟรีโซน จำกัด'
    assert first['invoiceNo'] == 'TEST-INV-01'
    assert first['invoiceDate'] == '07/08/2026'
    assert first['ins_tariffCode'] == '000087032290'
    assert first['thaiDescriptionOfGoods'].startswith('รถยนต์เก่า')
    assert first['endesc2'] == 'HR15(GASOLINE/1490),2011'
    assert first['productAttribute2'] == engines[0]
    assert first['productCode'] == '1490CC'
    assert first['netWeight'] == '1140'
    assert data[2]['originCountry'] == 'DE'               # MERCEDES BENZ
    assert data[2]['invoiceAmountForeign'] == '172700.5'
    assert data[2]['ins_tariffCode'] == 'XXXX'            # ไม่มีพิกัดในตาราง
    assert run['totals'] == {'package_amount': '3', 'package_unit': 'UN', 'gross_weight': '4330',
                             'price': '1022700.5', 'currency': 'JPY'}

    # อัปโหลดไฟล์ Excel ที่ได้ซ้ำ -> ใช้เลขเดิม ไม่สุ่มใหม่
    # (ไฟล์ทดสอบไม่มีสูตร จึงเติมค่าคอลัมน์ Q เองแทน Excel ที่คำนวณสูตรให้)
    out_xlsx = os.path.join(settings.output_dir, run['run_id'], run['excel'])
    wb = load_workbook(out_xlsx)
    for r, e in zip((20, 21, 22), engines):
        wb.active['Q%d' % r] = NC.add_number(wb.active['Q%d' % r].value, e.split('-')[1])
    wb.save(out_xlsx)
    again = pipeline.analyze(out_xlsx, settings, db)
    assert again.ok, again.errors
    assert [l.status for l in again.lines] == ['reuse'] * 3
    run2 = pipeline.commit(out_xlsx, settings, db)
    assert db.count() == 3 and [c['engine'] for c in run2['cars']] == engines


def test_block_layout(tmp_path, settings, db):
    src = make_jantrading(str(tmp_path / 'jan.xlsx'))
    run = pipeline.commit(src, settings, db)
    ws = load_workbook(os.path.join(settings.output_dir, run['run_id'], run['excel'])).worksheets[0]
    first = run['cars'][0]['engine']
    assert ws['D32'].value == '%s (GASOLINE/1490)' % first          # เขียนกลับในก้อนของรถคันแรก
    _, _, data = read_csv(run, settings)
    assert data[0]['endesc2'] == 'HR15(GASOLINE/1490),2011'


def test_full_number_not_in_db_is_kept_and_registered(tmp_path, settings, db):
    src = make_msfz(str(tmp_path / 'inv.xlsx'), engines=['HR15-7777777', '2GR', 'M271'])
    a = pipeline.analyze(src, settings, db)
    assert [l.status for l in a.lines] == ['keep', 'new', 'new']
    pipeline.commit(src, settings, db)
    assert db.find('7777777')[0]['CHASSIS'] == CARS[0][2]


def test_number_used_by_other_chassis_is_error(tmp_path, settings, db):
    with db.transaction() as tx:
        tx.add('7777777', 'HR15-7777777', 'SOMEONE-ELSE')
    src = make_msfz(str(tmp_path / 'inv.xlsx'), engines=['HR15-7777777', '2GR', 'M271'])
    a = pipeline.analyze(src, settings, db)
    assert any('ถูกใช้กับ chassis SOMEONE-ELSE' in e for e in a.errors)
    with pytest.raises(pipeline.AnalysisError):
        pipeline.commit(src, settings, db)
    assert db.count() == 1


def test_stale_summary_is_rejected(tmp_path, settings, db):
    src = make_msfz(str(tmp_path / 'inv.xlsx'))
    wb = load_workbook(src)
    wb.active['J20'] = 'MR20'          # แก้เลขเครื่องแต่ตารางสรุปยังเป็นค่าเก่า
    wb.save(src)
    a = pipeline.analyze(src, settings, db)
    assert any('ไม่ตรงกับเลขเครื่อง' in e for e in a.errors)


def test_failed_write_keeps_db_and_output_clean(tmp_path, settings, db, monkeypatch):
    src = make_msfz(str(tmp_path / 'inv.xlsx'))

    def broken(*a, **k):
        raise OSError('disk full')
    monkeypatch.setattr(pipeline, 'write_engines', broken)
    with pytest.raises(OSError):
        pipeline.commit(src, settings, db)
    assert db.count() == 0
    assert not os.listdir(settings.output_dir)


# ---------- template จริง (มีไฟล์เมื่อไหร่ก็รัน) ----------

@pytest.mark.parametrize('name, file_type, cars', [
    ('MASTER MSFZ.xlsx', 'MSFZ', 4),
    ('MASTER JTC.xlsx', 'JTC', 14),
    ('MASTER JANTRADING.xlsx', 'JAN TRADING', 15),
    ('MASTER HUAYO.xlsx', 'HUAYO', 2),
])
def test_real_templates(name, file_type, cars, settings, db):
    a = pipeline.analyze(fixture(name), settings, db)
    assert a.ok, a.errors
    assert a.invoice.file_type == file_type
    assert len(a.lines) == cars


def test_real_template_inconsistent_rows_reported(settings, db):
    a = pipeline.analyze(fixture('MASTER AICHI.xlsx'), settings, db)
    assert not a.ok


@pytest.mark.skipif(not excel_available(), reason='ต้องมี Microsoft Excel')
def test_excel_recalculation_matches_csv(settings, db):
    import dataclasses
    s = dataclasses.replace(settings, excel_method='excel')
    run = pipeline.commit(fixture('MASTER MSFZ.xlsx'), s, db)
    assert run['excel_method'] == 'excel'
    assert not [w for w in run['warnings'] if 'Excel คำนวณ' in w]


def test_cars_without_engine_cells_is_error(tmp_path, settings, db):
    src = make_msfz(str(tmp_path / 'inv.xlsx'))
    wb = load_workbook(src)
    for r in (20, 21, 22):
        wb.active['E%d' % r] = None        # ลบ chassis ในส่วน invoice -> หาช่องเลขเครื่องไม่เจอ
    wb.save(src)
    a = pipeline.analyze(src, settings, db)
    assert not a.ok and a.errors
