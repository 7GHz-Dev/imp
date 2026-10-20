# -*- coding: utf-8 -*-
"""worker.py กับ ERP-KOLA จำลอง + NETBAY จำลอง: รับงาน -> ตรวจ -> ออกเลข -> สร้างใบขน + Import CSV -> ส่ง Ref No. กลับ"""
import csv
import dataclasses
import io

import pytest
from openpyxl import Workbook, load_workbook

from conftest import CARS, make_msfz, serve
from mock_erp import KEY, create_mock_erp
from mock_netbay import create_mock
from worker import Worker

WITH_TARIFF = CARS[:2]        # คันที่ 3 ในชุดทดสอบไม่มีพิกัด


@pytest.fixture
def env(settings, db):
    erp, nb = create_mock_erp(), create_mock()
    erp_srv, erp_url = serve(erp)
    nb_srv, nb_url = serve(nb)
    s = dataclasses.replace(
        settings,
        netbay=dataclasses.replace(settings.netbay, url=nb_url, username='u1', password='p1', headless=True),
        erp=dataclasses.replace(settings.erp, url=erp_url, api_key=KEY, worker_name='imp@test', poll_seconds=1),
    )
    w = Worker(s, db=db)
    w.retry_waits = (0, 0)
    yield w, erp, nb.state
    erp_srv.shutdown()
    nb_srv.shutdown()


def engines(xlsx_bytes):
    ws = load_workbook(io.BytesIO(xlsx_bytes)).worksheets[0]
    return [ws['J%d' % r].value for r in (20, 21)]


def test_creates_draft_imports_csv_and_returns_ref_no(env, tmp_path, db):
    w, erp, netbay = env
    erp.add_task(make_msfz(str(tmp_path / 'FN.xlsx'), cars=WITH_TARIFF))
    w.run(once=True)

    assert erp.state['failed'] == []
    [done] = erp.state['completed']
    assert done['refNo'] == 'QENX000000099'
    assert done['url'] is None, 'URL ของ NETBAY พก UCODE ของ session — ห้ามส่งขึ้น ERP'
    assert done['fileName'] == 'TESTBL0001.xlsx'
    hr15, gr = engines(done['file'])
    assert hr15.startswith('HR15-') and gr.startswith('2GR-') and len(hr15) == len('HR15-1234567')

    assert netbay['created'] == 1 and netbay['decl']['bl'] == 'TESTBL0001'
    assert hr15.encode() in netbay['upload'][1]
    # จบทันทีที่ Import สำเร็จ: ไม่ Refresh ไม่รวมยอด ไม่คำนวณใหม่ ไม่ Save ไม่ Send to Customs
    assert (netbay['refresh'], netbay['popups'], netbay['recalc'], netbay['saved']) == (0, [], [], 0)
    assert 'sent_to_customs' not in netbay
    assert db.count() == 2

    beats = erp.state['heartbeats']
    assert all(b['mode'] == 'draft' and b['worker'] == 'imp@test' for b in beats)
    assert any(b.get('shot') for b in beats)
    assert any(b['busy'] == 'Draft ใบขน MU2026100001' for b in beats)
    assert beats[-1]['offline'] is True
    log = beats[-1]['log']
    assert any('Success!! Please check your invoice' in line for line in log)
    assert any('QENX000000099' in line for line in log)


def test_draft_again_for_same_bl_keeps_engine_numbers(env, tmp_path, db):
    w, erp, netbay = env
    src = make_msfz(str(tmp_path / 'FN.xlsx'), cars=WITH_TARIFF)
    erp.add_task(src)
    w.run(once=True)
    erp.add_task(src, job_id='JOB-0001')          # FAH ตีกลับ -> PAINT กดสร้าง Draft ใหม่ด้วยไฟล์เดิม
    w.run(once=True)

    first, second = erp.state['completed']
    assert engines(first['file']) == engines(second['file'])
    assert db.count() == 2                          # ไม่ออกเลขเพิ่ม
    assert netbay['created'] == 2


def test_result_is_resent_without_second_declaration(env, tmp_path):
    w, erp, netbay = env
    task = erp.add_task(make_msfz(str(tmp_path / 'FN.xlsx'), cars=WITH_TARIFF))
    erp.state['fail_complete'] = 2                   # ส่งผลกลับไม่ถึง ERP
    w.run(once=True)
    assert erp.state['completed'] == [] and netbay['created'] == 1

    task['status'] = 'QUEUED'                        # ERP ปล่อยงานที่ค้างให้หยิบใหม่
    w.run(once=True)
    [done] = erp.state['completed']
    assert done['refNo'] == 'QENX000000099'
    assert netbay['created'] == 1                    # ไม่สร้างใบขนซ้ำ


def test_invoice_problems_go_back_to_paint(env, tmp_path, db):
    w, erp, netbay = env
    blank = str(tmp_path / 'blank.xlsx')
    Workbook().save(blank)
    erp.add_task(blank)
    erp.add_task(make_msfz(str(tmp_path / 'other.xlsx'), cars=WITH_TARIFF), bl_no='OTHERBL999')
    pdf = tmp_path / 'FN.pdf'
    pdf.write_bytes(b'%PDF-1.4')
    erp.add_task(str(pdf))
    for _ in range(3):
        w.run(once=True)

    errors = [f['error'] for f in erp.state['failed']]
    assert 'ไม่รู้จักชนิดไฟล์' in errors[0]
    assert 'ไม่ตรงกับงาน (OTHERBL999)' in errors[1]
    assert 'ต้องเป็นไฟล์ Excel' in errors[2]
    assert erp.state['completed'] == []
    assert netbay['login'] is None and db.count() == 0   # ไม่แตะ NETBAY ไม่ออกเลข


def test_missing_tariff_goes_in_as_xxxx_like_before(env, tmp_path, db):
    w, erp, netbay = env
    erp.add_task(make_msfz(str(tmp_path / 'FN.xlsx')))   # คันที่ 3 (MERCEDES BENZ C200) ไม่มีพิกัด
    w.run(once=True)

    assert erp.state['failed'] == []
    [done] = erp.state['completed']
    assert done['refNo'] == 'QENX000000099' and netbay['created'] == 1
    head, _, *cars = csv.reader(io.StringIO(netbay['upload'][1].decode('tis-620')))
    tariff = head.index('ins_tariffCode')
    assert [c[tariff] for c in cars] == ['000087032290', '000087032449', 'XXXX']
    assert db.count() == 3
    assert any('พิกัดยังไม่ครบ 1 คัน' in line for line in erp.state['heartbeats'][-1]['log'])


def test_import_without_success_message_fails_with_what_netbay_said(env, tmp_path, monkeypatch):
    w, erp, netbay = env
    import core.netbay as nb
    monkeypatch.setattr(nb, 'IMPORT_TIMEOUT', 3)
    netbay['import_ok'] = False
    erp.add_task(make_msfz(str(tmp_path / 'FN.xlsx'), cars=WITH_TARIFF))
    w.run(once=True)
    [f] = erp.state['failed']
    assert 'success please check your invoice' in f['error']
    assert 'Error: invalid file format' in f['error']           # ข้อความบนหน้าจอ NETBAY ติดไปด้วย
    assert 'QENX000000099' in f['error']                         # บอกใบขนที่สร้างค้างไว้
    assert erp.state['completed'] == []
