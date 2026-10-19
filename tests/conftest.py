# -*- coding: utf-8 -*-
"""
ไฟล์ทดสอบ

synthetic invoice (ข้อมูลสมมติ) สร้างในเทสต์เอง — รันได้ทุกเครื่อง
tests/fixtures/*.xlsx (MASTER template จริง) ไม่อยู่ใน git — มีไฟล์เมื่อไหร่เทสต์ชุดนั้นจะรันเอง
"""
import dataclasses
import os
import shutil
import sys

import pytest
from openpyxl import Workbook

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

from core import config  # noqa: E402
from core.engine_db import EngineDB  # noqa: E402

FIXTURES = os.path.join(ROOT, 'tests', 'fixtures')

# (ยี่ห้อ, รุ่น, chassis, น้ำหนัก, ปี, เลขเครื่อง, cc, เชื้อเพลิง, ราคา)
CARS = [
    ('NISSAN', 'AD VAN', 'TEST12-000001', 1140, 2011, 'HR15', 1490, 'GASOLINE', 250000),
    ('TOYOTA', 'CROWN', 'TEST20-000002', 1690, 2008, '2GR', 3450, 'GASOLINE', 600000),
    ('MERCEDES BENZ', 'C200', 'WDDTEST0000000003', 1500, 2012, 'M271', 1790, 'GASOLINE', 172700.5),
]

TARIFF = [  # trf_num, trf_code, brand, model, fuel, cc_start, cc_end, w_start, w_end
    (1, 87032290, 'NISSAN', 'AD VAN', 'GASOLINE', 1001, 1500, 0, 99999),
    (2, 87032449, 'TOYOTA', 'CROWN', 'GASOLINE', 3001, 9999, 0, 99999),
]


def fixture(name):
    p = os.path.join(FIXTURES, name)
    if not os.path.isfile(p):
        pytest.skip('ไม่มีไฟล์ทดสอบ %s (ไม่อยู่ใน git)' % name)
    return p


def _summary(ws, cars, engines):
    """ตารางสรุปแถว 19+ แบบที่สูตรใน template คำนวณไว้ (เขียนเป็นค่า เหมือนค่าที่ Excel cache ไว้)"""
    heads = ['use_desc_good1', 'use_desc_good2', 'use_chassis_no', 'use_brand', 'use_weight', 'use_price', None,
             'car_no', 'brand', 'model', 'chassis_no', 'weight', 'year', 'cc', 'fuel', 'price', 'tariff']
    for j, h in enumerate(heads):
        if h:
            ws.cell(19, 16 + j, h)
    for i, ((brand, model, chassis, w, year, _, cc, fuel, price), eng) in enumerate(zip(cars, engines)):
        r = 20 + i
        values = ['USED %s %s,' % (brand, model), '%s(%s/%d),%d' % (eng.replace(' ', ''), fuel, cc, year),
                  chassis, brand, w, price, None, i + 1, brand, model, chassis, w, year, cc, fuel, price, 'XXXX']
        for j, v in enumerate(values):
            if v is not None:
                ws.cell(r, 16 + j, v)


def _header(ws, file_type):
    ws['P4'], ws['P5'], ws['P6'], ws['P14'] = file_type, '7/8/2026', 'TEST-INV-01', 'TESTBL0001'


def _tariff_sheet(wb):
    t = wb.create_sheet('datatariff')
    t.append([])
    t.append(['trf_num', 'trf_code', 'trf_brand', 'trf_model', 'trf_fuel',
              'trf_cc_start', 'trf_cc_end', 'trf_w_start', 'trf_w_end'])
    for row in TARIFF:
        t.append(list(row))


def make_msfz(path, cars=CARS, engines=None):
    """invoice แบบ MSFZ: 1 คันต่อแถว เลขเครื่องคอลัมน์ J"""
    engines = engines or [c[5] for c in cars]
    wb = Workbook()
    ws = wb.active
    ws.title = 'sheet1'
    _header(ws, 'MSFZ')
    for col, h in zip('ABCEHIJKLM', ['', 'MAKE', 'MODEL', 'CHASSISNO.', 'WEIGHT', 'YEAR', 'ENGINE NO', 'CC', 'FUEL', 'AMOUT']):
        ws['%s19' % col] = h or None
    for i, ((brand, model, chassis, w, year, _, cc, fuel, price), eng) in enumerate(zip(cars, engines)):
        r = 20 + i
        for col, v in zip('ABCEHIJKLM', [i + 1, brand, model, chassis, w, year, eng, cc, fuel, price]):
            ws['%s%d' % (col, r)] = v
    _summary(ws, cars, engines)
    _tariff_sheet(wb)
    wb.save(path)
    return path


def make_jantrading(path, cars=CARS):
    """invoice แบบ JAN TRADING: รถเป็นก้อน หัวข้อคอลัมน์ C ค่าคอลัมน์ D"""
    wb = Workbook()
    ws = wb.active
    _header(ws, 'JAN TRADING')
    engines = ['%s (%s/%d)' % (c[5], c[7], c[6]) for c in cars]
    r = 30
    for (brand, model, chassis, w, year, _, cc, fuel, price), eng in zip(cars, engines):
        ws['C%d' % r], ws['D%d' % r] = 'USED ' + brand, model
        ws['C%d' % (r + 1)], ws['D%d' % (r + 1)] = 'CHASSIS NO.', chassis
        ws['C%d' % (r + 2)], ws['D%d' % (r + 2)] = 'ENGINE NO.', eng
        ws['C%d' % (r + 3)], ws['D%d' % (r + 3)] = 'REG YEAR', year
        r += 4
    _summary(ws, cars, [c[5] for c in cars])
    _tariff_sheet(wb)
    wb.save(path)
    return path


@pytest.fixture
def settings(tmp_path):
    s = config.load(os.path.join(ROOT, 'config.example.ini'))
    data = tmp_path / 'data'
    shutil.copytree(os.path.join(ROOT, 'data'), data,
                    ignore=shutil.ignore_patterns('engine.db*'))
    return dataclasses.replace(
        s,
        db_path=str(data / 'engine.db'),
        output_dir=str(tmp_path / 'output'),
        work_dir=str(tmp_path / 'work'),
        csv_template=str(data / 'csv_template.xlsx'),
        brand_country=str(data / 'brand_country.xlsx'),
        tariff_file=str(data / 'DATA TARIFF.xlsx'),
        excel_method='openpyxl',
    )


@pytest.fixture
def db(settings):
    return EngineDB(settings.db_path)
